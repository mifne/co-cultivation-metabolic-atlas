"""Add subpanel interpretations and a linked market-price BOM to the report."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from revise_workstation_docx import (
    LIGHT_BLUE,
    NAVY,
    WHITE,
    add_hyperlink,
    find_paragraph,
    insert_paragraph_after,
    set_cell_margins,
    set_cell_shading,
    set_paragraph_text,
    set_repeat_table_header,
)


PRICE_DATE = "2026年8月28日"


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold


def add_price_source(cell, label: str, url: str | None, note: str) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if url:
        add_hyperlink(paragraph, label, url)
    else:
        paragraph.add_run(label)
    if note:
        paragraph.add_run("。" + note)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for index, (cell, width) in enumerate(zip(row.cells, widths)):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=82, start=105, bottom=82, end=105)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if index in (1, 2):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(7.8)


def add_panel_note(doc: Document, caption_prefix: str, clauses: list[tuple[str, str]]) -> None:
    anchor = find_paragraph(doc, caption_prefix)
    paragraph = insert_paragraph_after(anchor, style="Normal")
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.widow_control = True
    lead = paragraph.add_run("各パネルの読み取り　")
    lead.bold = True
    lead.font.color.rgb = RGBColor(38, 62, 88)
    for index, (label, text) in enumerate(clauses):
        if index:
            paragraph.add_run("　")
        marker = paragraph.add_run(label + " ")
        marker.bold = True
        paragraph.add_run(text)
    for run in paragraph.runs:
        run.font.size = Pt(8.3)
    paragraph.runs[0].font.size = Pt(8.5)


def build_budget_table(doc: Document) -> None:
    old = doc.tables[12]
    table = doc.add_table(rows=1, cols=4)
    table.style = old.style
    old._tbl.addprevious(table._tbl)
    old._tbl.getparent().remove(old._tbl)

    headers = ["区分・具体型番", "数量", "税込参考価格", "参照先・発注条件"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True)
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])

    rows = [
        (
            "GPU\nNVIDIA RTX PRO 4000 Blackwell\n900-5G147-2570-000",
            "3基",
            "359,980円/基\n小計 1,079,940円",
            "価格.com",
            "https://kakaku.com/item/K0001682019/",
            "掲載最安。国内正規品・同一SKU・在庫を確認",
        ),
        (
            "CPU\nAMD Ryzen Threadripper 9960X BOX",
            "1",
            "269,767円",
            "価格.com",
            "https://kakaku.com/item/K0001702503/",
            "国内正規流通と保証条件を確認",
        ),
        (
            "マザーボード\nASUS Pro WS TRX50-SAGE WIFI",
            "1",
            "175,050円",
            "価格.com",
            "https://kakaku.com/item/K0001597233/",
            "CEB、sTR5、PCIe x16スロット3本",
        ),
        (
            "128 GB ECC RDIMM（指定品）\nOWC4R8D52R4128P／64 GB×2",
            "1組",
            "757,790円",
            "Amazon.co.jp",
            "https://www.amazon.co.jp/dp/B0CLYYZZXB",
            "2026年8月25日確認値。ASUS QVL未掲載のため互換保証とMemTest必須",
        ),
        (
            "NVMe SSD\nWD_BLACK SN850X 2 TB\nWDS200T2X0E",
            "1",
            "56,980円",
            "価格.com",
            "https://kakaku.com/item/K0001467995/",
            "TLC、2 TB。swap 256 GBとscratch 300 GB以上を確保",
        ),
        (
            "電源\nCorsair HX1500i 2025\nCP-9020309-JP",
            "1",
            "43,800円",
            "価格.com",
            "https://kakaku.com/item/K0001697267/",
            "ATX 3.1、1500 W。標準12V-2x6は2本のため、3本目の純正対応ケーブルを業者が書面確認",
        ),
        (
            "ケース\nFractal Design Define 7 XL Solid\nFD-C-DEF7X-01",
            "1",
            "39,980円",
            "価格.com",
            "https://kakaku.com/item/K0001238045/",
            "SSI-CEB対応。3 GPU、電源、ケーブル曲げ半径を実機確認",
        ),
        (
            "CPUクーラー\nNoctua NH-U14S TR5-SP6",
            "1",
            "21,022円",
            "価格.com",
            "https://kakaku.com/item/K0001588511/",
            "sTR5対応。350 W連続負荷は温度試験で合否判定",
        ),
        (
            "追加ケースファン\nNoctua NF-A14 PWM",
            "3",
            "3,080円/個\n小計 9,240円",
            "価格.com",
            "https://kakaku.com/item/K0000787692/",
            "前面吸気・上面/背面排気は組立時に最適化",
        ),
        (
            "GPU支持具・純正電源ケーブル・小物予備費",
            "一式",
            "15,000円",
            "正式見積",
            None,
            "型番、必要本数、GPU側定格をBTO業者が確定",
        ),
        (
            "組立・BIOS・Ubuntu・3 GPU負荷試験・送料",
            "一式",
            "上限 80,000円",
            "業者見積",
            None,
            "研究ソフト導入は研究室側。受入ログを提出",
        ),
        (
            "本体・組立 税込参考総額",
            "—",
            "2,548,569円",
            "算定値",
            None,
            "上記の数量×参考価格。ポイント還元は控除しない",
        ),
        (
            "3年保守（別枠）",
            "一式",
            "上限 100,000円",
            "業者見積",
            None,
            "上限超過時は引取修理または別年度契約",
        ),
        (
            "保守込み上限見込み",
            "—",
            "2,648,569円",
            "算定値",
            None,
            "270万円枠に対して51,431円の余裕",
        ),
        (
            "条件付き減額候補\nG.SKILL Zeta R5 Neo 128 GB（32 GB×4）\nF5-6400R3239G32GQ4-ZR5NK",
            "1組",
            "183,980円\nOWC比 −573,810円",
            "パソコン工房",
            "https://www.pc-koubou.jp/products/detail.php?product_id=1213170",
            "掲載時は在庫切れ。4チャネルを使える一方、全4スロットを占有。QVL・互換保証・再入荷・MemTest確認時のみ採用",
        ),
    ]

    for row_index, row_data in enumerate(rows, start=1):
        cells = table.add_row().cells
        item, qty, price, label, url, note = row_data
        set_cell_text(cells[0], item, bold=row_index in (12, 14))
        set_cell_text(cells[1], qty)
        set_cell_text(cells[2], price, bold=row_index in (12, 14))
        add_price_source(cells[3], label, url, note)
        if row_index in (12, 14):
            for cell in cells:
                set_cell_shading(cell, LIGHT_BLUE)
        elif row_index == 15:
            for cell in cells:
                set_cell_shading(cell, "FFF2CC")

    set_table_geometry(table, [2550, 650, 1900, 4878])


def revise(source: Path, output: Path) -> None:
    doc = Document(str(source))

    add_panel_note(
        doc,
        "図1　",
        [
            ("(a)", "batch/GEMを8から128へ増やすとGPU使用率は56.9%から98.3%へ上がり、十分な同時候補がGPUを飽和させる。"),
            ("(b)", "VRAMは0.6から7.7 GiBへ増え、処理量はbatch 64で最大化した後に128で低下するため、容量上限直前が必ずしも最高効率ではない。"),
            ("(c)", "CUDA化でCPU使用率は低下するが、32環境では84.2%まで戻りRSSも4.61 GiBとなるため、GPUへ仕事を供給するCPU/RAMも増強対象である。"),
            ("(d)", "提案機はCPU thread 3倍、RAM 2.7倍、GPU当たりVRAM 3倍、総VRAM 9倍であり、3つの独立shardを同時実行する容量を確保する。"),
        ],
    )
    add_panel_note(
        doc,
        "図2　",
        [
            ("(A)", "同時環境数の増加に伴いGPU使用率が上昇し、concurrent CUDA streamsは全点で逐次実行を上回るため、複数環境の束ね方が利用率を決める。"),
            ("(B)", "有効な3-GEM FBA候補処理量はbatch 64付近で最大となり、128では急落するため、本番初期値は64とする。"),
            ("(C)", "VRAMはbatchにほぼ単調増加し、128でRTX 4060の8 GB上限に接近することが、(B)の処理量低下と整合する。"),
            ("(D)", "16環境・960 transitionではCPU/HiGHS 12.2に対してCUDA surrogate 85.8 transition/s、7.01倍である。24-step比較の4.99倍とは試験条件が異なるため別の実測値として扱う。"),
        ],
    )
    add_panel_note(
        doc,
        "図3　",
        [
            ("(a)", "現行wall timeはGPU推論41.8%、環境更新等58.2%であり、GPUだけを速めてもCPU側を残すとAmdahl則で上限が生じる。"),
            ("(b)", "GPU段は1/2/3基で現行比3.20/4.18/5.94倍と予測し、並列効率損失を含めて線形3倍より保守的に置いた。"),
            ("(c)", "独立16環境shardを1/2/3本へ増やすとend-to-end処理量は111/211/300 transition/sとなり、3基で現行比約3.50倍を見込む。"),
            ("(d)", "生産カーネル容量は12.3/23.3/33.2千 prediction/sへ増え、FBA候補生成段には十分な余裕を持つ。"),
            ("(e)", "固定仕事時間はend-to-endで28.6%、GPU段のみで16.8%まで短縮するため、残差が環境更新・転送・同期の改善対象となる。"),
            ("(f)", "総VRAM 9倍は共有メモリではなく3基合計であり、1つの72 GBモデルではなく、seed・agent・環境群をGPUごとに分離する容量である。"),
        ],
    )

    build_budget_table(doc)

    set_paragraph_text(
        find_paragraph(doc, "約270万円に収める調達案"),
        "市場価格参照による調達案　2026年8月28日の価格.com掲載最安値等を数量別に積み上げると、指定OWC 128 GBを含む本体・組立は2,548,569円、"
        "3年保守込みは2,648,569円である。270万円枠に対する余裕は51,431円にとどまるため、価格変動、送料、在庫、純正GPU電源ケーブルの追加で超過し得る。"
        "正式発注では同一型番、国内保証、納期、3 GPU配線を含む一式見積で再確認する。",
        bold_prefix="市場価格参照による調達案　",
    )
    set_paragraph_text(
        find_paragraph(doc, "見積は『"),
        "価格の扱い　価格.com掲載値は2026年8月28日閲覧の税込参考最安値で、送料・ポイント・在庫変動を含む保証額ではない。"
        "OWCメモリはAmazon.co.jpで2026年8月25日に確認した757,790円を使用した。価格比較ページと商品ページへのリンクを表内に付し、正式BOMでは販売店、在庫、保証、納期を再確認する。",
        bold_prefix="価格の扱い　",
    )
    set_paragraph_text(
        find_paragraph(doc, "第2 NVMeを"),
        "標準案は第2 NVMeを将来増設とし、2 TB 1基に限定する。これ以上SSD容量を削ると256 GB swapと300 GB以上のscratchを同時確保できないため、減額対象にしない。",
    )
    set_paragraph_text(
        find_paragraph(doc, "解析ソフトの導入"),
        "第一の減額候補はメモリ調達方法である。G.SKILL Zeta R5 Neo 128 GB（32 GB×4）の掲載価格183,980円を採用できればOWC比573,810円を削減できるが、"
        "掲載時点では在庫切れであり、4スロットを全て占有する。ASUS QVL、販売店の互換保証、再入荷、MemTest合格を満たす場合だけ代替し、条件を満たさなければ指定OWCを維持する。",
    )
    set_paragraph_text(
        find_paragraph(doc, "3年オンサイト保守"),
        "第二の減額候補は保守の年度分離である。3年保守10万円を別年度契約とすれば当年度支出を10万円下げられる。解析ソフト導入は研究室側で行い、業者作業はUbuntu、driver、CUDA、3 GPU負荷試験までに限定する。",
    )
    set_paragraph_text(
        find_paragraph(doc, "さらに減額が必要"),
        "減額しても、非SFF GPU 3基、Threadripper/TRX50、128 GB ECCという研究要件と、1,500 W電源・冷却・3 GPU認識試験は維持する。"
        "ケース小型化、非純正電源ケーブル、負荷試験の省略は、干渉・過熱・電源障害のリスクを増やすため減額対象としない。",
    )

    executive = find_paragraph(doc, "導入判断")
    set_paragraph_text(
        executive,
        executive.text.replace("本体・組立は約267.7万円、3年保守は別枠上限10万円とする。", "市場価格参照では本体・組立254万8,569円、3年保守込み264万8,569円を見込む。"),
        bold_prefix="導入判断　",
    )
    summary = find_paragraph(doc, "依頼文")
    set_paragraph_text(
        summary,
        summary.text.replace("本体・組立は約267.7万円、3年保守は別枠上限10万円とする。", "本体・組立の税込参考総額は254万8,569円、3年保守込みは264万8,569円とする。"),
        bold_prefix="依頼文　",
    )

    props = doc.core_properties
    props.comments = "2026-08-28 revision: per-panel interpretations and linked market-price BOM added."

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    revise(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
