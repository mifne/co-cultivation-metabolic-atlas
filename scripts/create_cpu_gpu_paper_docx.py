#!/usr/bin/env python3
"""Build a Japanese manuscript-style DOCX for CPU-to-GPU dFBA migration."""

from __future__ import annotations

import json
from pathlib import Path
import statistics

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/cpu_gpu_paper_20260902"
FIG = OUT / "figures"
DOCX = OUT / "CPU_GPU_dFBA_manuscript_ja_PHA_accuracy_added.docx"
VALIDATION = ROOT / "models/cooperative_surrogate/cooperative_neural_reranker_34448_dagger2_rebased.validation.json"
BASELINE_VALIDATION = ROOT / "models/cooperative_surrogate/cooperative_neural_reranker_33728_uniform_late_trained_pha256.validation.json"
DAGGER1_VALIDATION = ROOT / "results/pha_validation_dagger_rebased_5seed.json"
CALIBRATION = ROOT / "results/pha_hyperparameter_calibration_3seed.json"
PROFILE = ROOT / "results/cooperative_gpu_qp_only_33728_trained_k128_24h_seed20260901.json"
SCALING = ROOT / "results/gpu_qp_batch_scaling_rtx4060.json"
SMOKE = ROOT / "results/gpu_qp_parallel_training_smoke_dagger2.json"

INK = RGBColor(31, 41, 51)
BLUE = RGBColor(0, 114, 178)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(100, 112, 124)
LIGHT_FILL = "F2F4F7"
BLUE_FILL = "E4EFF7"
GREEN_FILL = "E2F1EB"
CAUTION_FILL = "FBE8DD"


def set_font(run, name="IPAexMincho", size=10.5, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.40)

    normal = doc.styles["Normal"]
    normal.font.name = "IPAexMincho"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "IPAexMincho")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.18

    for name, size, before, after, color in (
        ("Heading 1", 14, 14, 6, BLUE),
        ("Heading 2", 11.5, 10, 4, DARK_BLUE),
        ("Heading 3", 10.5, 8, 3, DARK_BLUE),
    ):
        style = doc.styles[name]
        style.font.name = "IPAexGothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "IPAexGothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "IPAexMincho"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "IPAexMincho")
    caption.font.size = Pt(8.7)
    caption.font.color.rgb = MUTED
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.05

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    set_font(hp.add_run("CPU-to-GPU cooperative dFBA | Technical manuscript"), "IPAexGothic", 8, color=MUTED)
    paragraph_bottom_border(hp, "BFC7CE", 4)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(fp.add_run("2026-09-02  |  "), "IPAexGothic", 8, color=MUTED)
    add_field(fp, "PAGE")


def paragraph_bottom_border(paragraph, color, size):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    set_font(run, "IPAexGothic", 8, color=MUTED)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0072B2")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend((color, underline))
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend((rpr, text_node))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_body(doc, text: str, *, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead:
        set_font(p.add_run(bold_lead), "IPAexGothic", 10.5, bold=True)
    set_font(p.add_run(text), "IPAexMincho", 10.5)
    return p


def add_equation(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(text), "Cambria Math", 10.5, italic=True)
    return p


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths_dxa)))
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), "120")
    tblind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths_dxa[index]))
            tcw.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def fill_table(table, headers: list[str], rows: list[list[str]], widths: list[int], font_size=8.5):
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), "IPAexGothic", font_size, bold=True)
        set_cell_shading(cell, BLUE_FILL)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = ""
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(value), "IPAexMincho", font_size)
    set_table_geometry(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def add_figure(doc, filename: str, width: float, caption: str, alt: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline = run.add_picture(str(FIG / filename), width=Inches(width))
    inline._inline.docPr.set("descr", alt)
    cp = doc.add_paragraph(caption, style="Caption")
    return cp


def add_reference(doc, number: int, citation: str, url: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    set_font(p.add_run(f"[{number}] {citation} "), "IPAexMincho", 8.5)
    add_hyperlink(p, "DOI / source", url)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    validation = json.loads(VALIDATION.read_text(encoding="utf-8"))
    baseline = json.loads(BASELINE_VALIDATION.read_text(encoding="utf-8"))
    dagger1 = json.loads(DAGGER1_VALIDATION.read_text(encoding="utf-8"))
    calibration = json.loads(CALIBRATION.read_text(encoding="utf-8"))
    profile = json.loads(PROFILE.read_text(encoding="utf-8"))
    scaling = json.loads(SCALING.read_text(encoding="utf-8"))
    smoke = json.loads(SMOKE.read_text(encoding="utf-8"))
    runs = validation["runs"]
    cpu_mean = statistics.mean(row["exact_seconds"] for row in runs)
    gpu_mean = statistics.mean(row["surrogate_seconds"] for row in runs)
    scaling_rows = scaling["rows"]
    throughput_scale = scaling_rows[-1]["environments_per_second"] / scaling_rows[0]["environments_per_second"]
    latency_drop = 100 * (1 - scaling_rows[-1]["per_environment_milliseconds"] / scaling_rows[0]["per_environment_milliseconds"])
    pha_mean_reduction = 100 * (
        1 - validation["summary"]["pha_relative_error_mean"] / baseline["summary"]["pha_relative_error_mean"]
    )
    pha_max_reduction = 100 * (
        1 - validation["summary"]["pha_relative_error_max"] / baseline["summary"]["pha_relative_error_max"]
    )
    profile_gpu = profile["surrogate"]["solver"]
    qp_per_step = profile_gpu["cooperative"]["gpu_qp_seconds"] / profile["steps"]
    search_per_step = profile_gpu["cooperative"]["build_seconds"]
    qp_share = 100 * qp_per_step / profile_gpu["mean_step_timing"]["total_seconds"]
    search_share = 100 * search_per_step / profile_gpu["mean_step_timing"]["total_seconds"]

    doc = Document()
    style_document(doc)
    doc.core_properties.title = "三種GEM共同dFBAのCPU–GPU移行"
    doc.core_properties.subject = "GPU batch-QP acceleration for cooperative dFBA reinforcement learning"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "dFBA, GEM, GPU, QP, reinforcement learning, microbial coculture"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(7)
    set_font(
        title.add_run("三種GEM共同dFBA強化学習における\nCPU線形計画からGPUバッチQPへの移行"),
        "IPAexGothic",
        18,
        bold=True,
        color=DARK_BLUE,
    )
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    set_font(
        subtitle.add_run("Implementation, performance, scalability, and novelty assessment"),
        "Arial",
        10.5,
        italic=True,
        color=MUTED,
    )
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(14)
    set_font(meta.add_run("Technical manuscript  |  2026-09-02  |  RTX 4060 Laptop GPU"), "IPAexGothic", 8.8, color=MUTED)

    doc.add_heading("要旨", level=1)
    abstract = (
        "微生物3種のgenome-scale metabolic models（GEMs）を統合した共同dynamic flux balance analysis（dFBA）では、"
        "環境stepごとに大規模線形計画（LP）を反復するため、強化学習の試行回数が計算時間に直結する。"
        "本研究では、CPU HiGHSによる三段階共同LPを参照系とし、34,448個の厳密FBAフラックスをアンカーとする"
        "ニューラル順位付け、アンカー凸包上のGPUバッチ二次計画（QP）投影、およびPPO複数環境を集約する"
        "単一CUDAマイクロバッチサービスを実装した。GPU自由rolloutが訪問する状態の分布ずれを補正するため、"
        "2反復のDAgger型exact再ラベルにより720アンカーを追加した。5個の独立held-out seed、各120 stepの同一action比較では、CPU参照系"
        f"{cpu_mean:.2f} sに対しGPU系{gpu_mean:.2f} sであり、平均{validation['summary']['speedup_mean']:.2f}倍"
        f"（95% CI {validation['summary']['speedup_95_ci'][0]:.2f}–{validation['summary']['speedup_95_ci'][1]:.2f}）の高速化を得た。"
        f"GPUカーネルのスループットはbatch 1から64で{throughput_scale:.2f}倍となり、環境当たり遅延は{latency_drop:.1f}%減少した。"
        "4環境×128 stepのPPO様smoke testでは512/512要求をGPU上で有効解へ変換し、CPU LP呼出しは0回であった。"
        f"PHA終点相対誤差はbaselineの平均{100*baseline['summary']['pha_relative_error_mean']:.2f}%、最大"
        f"{100*baseline['summary']['pha_relative_error_max']:.2f}%から、平均{100*validation['summary']['pha_relative_error_mean']:.2f}%、最大"
        f"{100*validation['summary']['pha_relative_error_max']:.2f}%へ低下し、事前設定した最大1%の同等性基準を全seedで満たした。"
        "したがって、本実装は探索用GPUバックエンドとして有効であるが、最終的な科学評価にはCPU厳密解による再評価を要する。"
    )
    add_body(doc, abstract)
    key = doc.add_paragraph()
    key.paragraph_format.space_after = Pt(12)
    set_font(key.add_run("キーワード："), "IPAexGothic", 9.5, bold=True)
    set_font(key.add_run("共同dFBA；ゲノムスケール代謝モデル；GPU；バッチQP；PPO；微生物共培養"), "IPAexMincho", 9.5)

    doc.add_heading("1. 緒言", level=1)
    add_body(
        doc,
        "FBAは定常状態条件 Sv = 0、反応上下限および生物学的目的関数をLPとして解くことで代謝表現型を推定する[1]。"
        "dFBAはこのLPを外部培養液の動的物質収支へ埋め込むため、右辺評価のたびに一つ以上のLP解を必要とする[2,3]。"
        "共同培養では複数GEMと共有培地制約が加わり、さらにPPO等の強化学習では多数の環境・seed・episodeを評価するため、"
        "反復LPが主要な計算負荷となる。"
    )
    add_body(
        doc,
        "微生物共培養への強化学習適用[4]、GEMをエージェントとして扱うSPAM-DFBA[5]、"
        "ニューラル・機構論ハイブリッドGEM[6]、GPU上のバッチLP/QP[7,8]はそれぞれ報告されている。"
        "しかし、三種GEMの共有培地共同dFBAを対象に、厳密FBA解の凸包で質量収支を保持しながら、"
        "GPUバッチQPとPPO複数環境サービスを統合した構成は、今回確認した文献範囲では見いだせなかった。"
        "本稿では実装、実測性能、スケーラビリティ、精度限界および新規性候補を整理する。"
    )

    doc.add_heading("2. 材料および方法", level=1)
    doc.add_heading("2.1 三種GEM共同dFBA", level=2)
    add_body(
        doc,
        "対象系はActinoplanes sp. OR16、Rhizobacter gummiphilus NS21、Lactobacillus plantarum WCFS1の3モデルからなる。"
        "反応総数6,585、代謝物総数4,429、遺伝子総数3,899であり、block-diagonal stoichiometric matrixを用いて各種の内部定常状態を同時に課した。"
    )
    table = doc.add_table(rows=1, cols=5)
    fill_table(
        table,
        ["GEM", "反応", "代謝物", "遺伝子", "交換反応"],
        [
            ["Actinoplanes sp. OR16", "2,584", "1,668", "1,596", "284"],
            ["R. gummiphilus NS21", "2,792", "1,785", "1,390", "332"],
            ["L. plantarum WCFS1", "1,209", "976", "913", "145"],
            ["合計", "6,585", "4,429", "3,899", "761"],
        ],
        [3000, 1450, 1450, 1450, 1530],
    )
    doc.add_paragraph("表1　三種GEMの計算規模。", style="Caption")

    doc.add_heading("2.2 CPU参照アルゴリズム", level=2)
    add_body(
        doc,
        "CPU参照系はSciPy/HiGHSを用いる共同LPである。第1段階では各種増殖率の最小値 g を最大化し、"
        "共存可能な共通増殖下限を決定する。第2段階はlive objectiveを使用する設定に備えるが、本評価では無効とした。"
        "第3段階ではgの99%を保持しつつ、バイオマス重み付き環境交換フラックスの絶対値和を最小化した。"
        "このLPをdt = 0.2 h、120 step（24 h相当）の各時点で再構築・求解した。"
    )
    add_equation(doc, "maximize  g    subject to    S v = 0,   l ≤ v ≤ u,   g ≤ μs,   Ashared(X)v ≤ C/Δt")

    doc.add_heading("2.3 GPUニューラル順位付けと凸包QP投影", level=2)
    add_body(
        doc,
        "オフラインでCPU厳密共同LPから32,768候補を収集し、独立した境界軌道960候補を追加して33,728アンカーのbaseline辞書を構築した。"
        "ニューラルネットワークは97個の変動context特徴から、増殖・PHA・交換反応を含む212 decision fluxを推定し、"
        "GPU常駐辞書を順位付けする。通常は上位128アンカーを選ぶ。"
    )
    add_body(
        doc,
        "各アンカー vi はオフライン厳密解であるため Svi = 0 を満たす。候補行列Vと単体重みwを用いて v = Vᵀw とすれば、"
        "w ≥ 0かつ1ᵀw = 1の下で Sv = 0が構成上保持される。GPU QPは先頭候補への距離を最小化しながら、"
        "その時点の反応上下限と共有培地制約へ投影する。"
    )
    add_equation(doc, "minw  ½‖w − e0‖²    s.t.   w ≥ 0,  1ᵀw = 1,  l ≤ Vᵀw ≤ u,  AsharedVᵀw ≤ b")
    add_body(
        doc,
        "128候補で有効解が得られない場合のみ2,048候補へ拡張し、各種上位6候補のblock compositionを追加して最大1,200回の"
        "primal–dual反復を行う。この再試行もCUDA上で完結し、CPU LPへはフォールバックしない。"
    )

    doc.add_heading("2.4 PPO複数環境のCUDAマイクロバッチ化", level=2)
    add_body(
        doc,
        "各PPO workerが約1 GiBの辞書とCUDA contextを個別保持すると、VRAM重複と小規模kernel launchが生じる。"
        "そこでspawnされた単一GPU service processが辞書、順位付け器、QP projectorを所有し、2 msのwindow内に到着した"
        "環境要求を最大64件までまとめる。workerはflux結果のみを受け取り、dFBA状態積分を継続する。"
    )
    add_figure(
        doc,
        "Figure_1_CPU_GPU_architecture.png",
        6.55,
        "図1　CPU参照経路とGPUオンライン経路。（a）CPUでは各stepで6,585変数規模の共同LPを構築し、共存増殖とparsimonious exchangeを逐次求解する。（b）GPUでは複数RL環境を単一CUDA serviceへ集約し、ニューラル順位付け、128アンカーの高速feasibility scan、凸包QP投影を行う。通常候補で不成立の要求のみ2,048アンカーへ拡張する。",
        "CPU reference LP path and GPU neural-ranking, adaptive anchor retry, and batched QP path.",
    )

    doc.add_heading("2.5 ベンチマーク設計", level=2)
    add_body(
        doc,
        "性能比較はRTX 4060 Laptop GPU搭載環境で行った。5個の独立action seedについて同一120-step action列をCPU参照系とGPU系へ入力した。"
        "speed-upはCPU実行時間/GPU実行時間とし、平均の95%信頼区間はt分布で算出した。スケーリング試験は同一GPU上でbatch size 1、8、32、64を各3反復した。"
        "並列サービス試験は4環境×128 stepのPPO様actionを用いた。"
    )

    doc.add_heading("2.6 PHA終点精度改善実験", level=2)
    add_body(
        doc,
        "baseline辞書はCPU exact軌道から構築されるため、GPU近似器自身が生成した自由rolloutでは学習時と異なる状態へ遷移し、"
        "一stepの小さなフラックス差が外部培地、菌体量、PHA蓄積量を介して終点誤差へ累積する可能性がある。"
        "これは逐次予測におけるexposure biasに相当する。そこで、学習器が誘導した状態をexpertで再ラベルするDAggerの考え方[11]を、"
        "機構論的dFBAへ適用した。分布ずれを扱う点はscheduled sampling[12]とも共通する。"
    )
    add_body(
        doc,
        "各反復では3本のGPU-only 120-step軌道からliveな97特徴contextを収集し、同一の反応上下限、菌体量、共有基質供給量を用いて"
        "CPU HiGHS三段階共同LPをofflineで再求解した。得られたexact-feasible fluxを各反復360件、計720件追加し、"
        "候補数を33,728から34,448へ増加させた。既存ニューラル重みは汎化性能を維持するため再学習せず、新辞書へrebaseした。"
        "補正用seedと最終validation seed（20260901–20260905）は重複させなかった。"
    )
    add_equation(doc, "terminal PHA relative error = |PHA_GPU(T) − PHA_CPU(T)| / PHA_CPU(T)")
    add_body(
        doc,
        "主要評価項目は5 seedにおけるPHA終点相対誤差の平均および最大値とした。適格条件は、最大PHA誤差≤1%、"
        "GPU解受理率≥90%、最大菌体量絶対誤差≤0.01 g/L、平均speed-up≥1と事前定義した。"
        "性能の95%信頼区間はseed別speed-upのt分布区間であり、n = 5のため精度値は記述統計として扱った。"
    )
    add_figure(
        doc,
        "Figure_2_DAgger_rollout_relabeling.png",
        6.45,
        "図2　GPU rollout状態のDAgger型exact再ラベル。（a）CPU exact軌道のみで構築した辞書とGPU自由rolloutの状態分布がずれることで、実行可能性を保ったまま終点PHA誤差が累積する。（b）GPU訪問状態を収集し、同じlive制約でCPU HiGHSを再求解してexact-feasible anchorを追加する工程。各反復360件、2反復、計720件を追加し、validation seedは補正に使用しなかった。",
        "Distribution shift in free GPU rollout and the two-iteration DAgger-style exact relabeling workflow.",
    )

    doc.add_page_break()
    doc.add_heading("3. 結果", level=1)
    doc.add_heading("3.1 24時間軌道の実行時間", level=2)
    add_body(
        doc,
        f"CPU参照系の平均実行時間は{cpu_mean:.2f} s、GPU系は{gpu_mean:.2f} sであった。seed別speed-upは"
        f"{min(row['speedup'] for row in runs):.2f}–{max(row['speedup'] for row in runs):.2f}倍、平均"
        f"{validation['summary']['speedup_mean']:.2f}倍（95% CI {validation['summary']['speedup_95_ci'][0]:.2f}–"
        f"{validation['summary']['speedup_95_ci'][1]:.2f}）であり、全seedで1倍を上回った。"
    )
    add_figure(
        doc,
        "Figure_3_performance_and_accuracy.png",
        6.45,
        "図3　CPU–GPU性能比較と終点精度。（a）同一action列による120-step軌道の実行時間。（b）seed別speed-upと平均値、網掛けは平均の95%信頼区間。（c）PHA終点相対誤差。破線は事前設定した最大1%の同等性基準である。n = 5。",
        "Runtime bars, seed-level speed-up with confidence interval, and terminal PHA relative error.",
    )

    doc.add_heading("3.2 バッチスケーラビリティ", level=2)
    add_body(
        doc,
        f"GPU kernel throughputはbatch 1の{scaling_rows[0]['environments_per_second']:.1f} environments/sからbatch 64の"
        f"{scaling_rows[-1]['environments_per_second']:.1f} environments/sへ増加し、{throughput_scale:.2f}倍となった。"
        f"環境当たり遅延は{scaling_rows[0]['per_environment_milliseconds']:.3f} msから"
        f"{scaling_rows[-1]['per_environment_milliseconds']:.3f} msへ{latency_drop:.1f}%低下した。"
        f"一方、peak allocated VRAMは{scaling_rows[0]['gpu_peak_allocated_mib']/1024:.2f} GiBから"
        f"{scaling_rows[-1]['gpu_peak_allocated_mib']/1024:.2f} GiBへ増加した。"
    )
    add_figure(
        doc,
        "Figure_4_GPU_batch_scaling.png",
        6.4,
        "図4　RTX 4060 Laptop GPUにおけるbatch scaling。（a）GPU順位付け＋投影kernelのthroughput。（b）環境当たりlatency。（c）PyTorch peak allocated VRAM。各点は3反復の平均、全条件でfeasible fraction = 1.0。",
        "GPU throughput, per-environment latency, and peak allocated VRAM across batch sizes.",
    )

    doc.add_heading("3.3 並列サービスの有効解回復", level=2)
    add_body(
        doc,
        f"4環境×128 stepでは{smoke['service']['requests']}要求を{smoke['service']['batches']} microbatchとして処理した。"
        f"平均batch sizeは{smoke['service']['mean_batch_size']:.2f}、最大は{smoke['service']['max_observed_batch']}であった。"
        f"通常128候補で{smoke['service']['requests']-smoke['service']['retry_requests']}要求が成立し、残る"
        f"{smoke['service']['retry_requests']}要求は2,048候補のGPU再試行で全件回復した。最終feasible rateは100%、CPU LP呼出しは0回であった。"
    )
    add_figure(
        doc,
        "Figure_5_parallel_service_robustness.png",
        6.25,
        "図5　PPO様4環境smoke testにおけるGPU service。（a）観測されたmicrobatch size分布。（b）512要求の処理フロー。507要求は通常候補で成立し、5要求はGPU候補拡張で回復した。",
        "Observed microbatch histogram and adaptive GPU retry recovery flow.",
    )

    doc.add_heading("3.4 科学的同等性", level=2)
    add_body(
        doc,
        f"600/600 stepでGPU有効解が得られ、最大バイオマス終点絶対誤差は"
        f"{validation['summary']['max_biomass_absolute_error_g_l']:.6f} g/Lであった。PHA終点相対誤差は平均"
        f"{100*validation['summary']['pha_relative_error_mean']:.2f}%、最大{100*validation['summary']['pha_relative_error_max']:.2f}%であった。"
        "したがって、物理feasibility、終点精度、速度の全事前条件を満たし、validation statusはqualifiedとなった。"
    )

    table2 = doc.add_table(rows=1, cols=3)
    fill_table(
        table2,
        ["評価項目", "実測値", "解釈"],
        [
            ["120-step speed-up", f"{validation['summary']['speedup_mean']:.2f}× (95% CI {validation['summary']['speedup_95_ci'][0]:.2f}–{validation['summary']['speedup_95_ci'][1]:.2f})", "全seedで高速化"],
            ["GPU有効解", "600/600 step", "CPU LP不要"],
            ["PHA終点誤差", f"平均 {100*validation['summary']['pha_relative_error_mean']:.2f}%; 最大 {100*validation['summary']['pha_relative_error_max']:.2f}%", "全seedで1%以下"],
            ["Batch 1→64", f"Throughput {throughput_scale:.2f}×", "環境並列が有効"],
            ["PPO様service", "512/512 valid", "5件をGPU再試行で回復"],
        ],
        [2800, 3300, 3260],
        font_size=8.2,
    )
    doc.add_paragraph("表2　主要な性能・精度指標。", style="Caption")

    doc.add_heading("3.5 PHA終点誤差のseed別改善", level=2)
    add_body(
        doc,
        f"baselineのPHA終点相対誤差は平均{100*baseline['summary']['pha_relative_error_mean']:.3f}%、最大"
        f"{100*baseline['summary']['pha_relative_error_max']:.3f}%であった。2反復補正後は平均"
        f"{100*validation['summary']['pha_relative_error_mean']:.3f}%、最大{100*validation['summary']['pha_relative_error_max']:.3f}%となり、"
        f"平均誤差を{pha_mean_reduction:.1f}%、最大誤差を{pha_max_reduction:.1f}%低減した。5 seedすべてで1%未満となり、"
        "CPU exactとGPU終点PHAは0.0510–0.0540 g/Lの範囲で近接した。"
    )
    add_figure(
        doc,
        "Figure_6_PHA_endpoint_accuracy.png",
        6.35,
        "図6　PHA終点精度の詳細。（a）同一held-out seedにおけるbaselineとDAgger-2の対応比較。数字はseed番号、破線は1%基準。（b）CPU exactとGPUの終点PHA、破線はy = x。（c）DAgger-2のPHA終点絶対誤差。（d）平均および最大相対誤差のbefore–after比較。n = 5。",
        "Paired terminal PHA errors, CPU-GPU parity, absolute errors, and aggregate before-after accuracy.",
    )
    seed_table = doc.add_table(rows=1, cols=6)
    fill_table(
        seed_table,
        ["Seed", "CPU PHA\n(g/L)", "GPU PHA\n(g/L)", "相対誤差\n(%)", "Speed-up", "GPU解"],
        [
            [
                str(row["seed"]),
                f"{row['exact_pha_g_l']:.6f}",
                f"{row['gpu_pha_g_l']:.6f}",
                f"{100*row['pha_relative_error']:.3f}",
                f"{row['speedup']:.2f}×",
                f"{row['gpu_qp_accepts']}/{validation['steps']}",
            ]
            for row in runs
        ],
        [1600, 1600, 1600, 1500, 1450, 1610],
        font_size=7.7,
    )
    doc.add_paragraph("表3　held-out seed別のPHA終点値、相対誤差、速度およびGPU有効解数。", style="Caption")

    doc.add_heading("3.6 DAgger反復による精度・速度推移", level=2)
    add_body(
        doc,
        f"第1反復では候補数34,088、平均誤差{100*dagger1['summary']['pha_relative_error_mean']:.3f}%、最大"
        f"{100*dagger1['summary']['pha_relative_error_max']:.3f}%であり、最大1%基準には未達であった。第2反復で候補数34,448とすると"
        f"最大誤差は{100*validation['summary']['pha_relative_error_max']:.3f}%まで低下した。候補数はbaseline比2.14%増に留まり、"
        f"最終speed-upは{validation['summary']['speedup_mean']:.2f}倍を維持した。"
    )
    add_figure(
        doc,
        "Figure_7_DAgger_iteration_progress.png",
        6.35,
        "図7　DAgger反復に伴う性能推移。（a）PHA終点相対誤差の平均と最大値。（b）baselineから追加したexact anchor数と総候補数。（c）各段階の平均speed-up。DAgger-1では平均がほぼ1%まで低下したが最大値は未達であり、第2反復で最大1%基準を満たした。",
        "PHA error, exact-anchor count, and speed-up across baseline and two DAgger iterations.",
    )

    doc.add_heading("3.7 実行時ボトルネック", level=2)
    add_body(
        doc,
        f"profile軌道ではCPU HiGHSの平均step時間0.589 sに対し、GPU系は0.199 sであった。GPU QP投影時間は平均"
        f"{1000*qp_per_step:.2f} ms/step、全step時間の{qp_share:.1f}%に留まった。一方、辞書検索・候補materializationは"
        f"{1000*search_per_step:.2f} ms/step、全step時間の{search_share:.1f}%を占めた。QP反復は大部分のstepで0回で終了しており、"
        "現在の律速はQP収束ではなく、34,448×97 context距離評価、候補tensor生成、およびPython dFBA状態更新であった。"
    )
    add_figure(
        doc,
        "Figure_8_runtime_bottleneck_profile.png",
        6.35,
        "図8　120-step profileにおける実行時間内訳。（a）CPU HiGHSとGPU surrogateの平均dFBA step時間。（b）GPU step内訳。辞書検索・候補materializationが53.5%を占め、QP投影は4.4%であった。単一軌道による機構別profileであり、seed間の統計比較ではない。",
        "CPU-GPU mean step timing and detailed GPU bottleneck decomposition.",
    )

    doc.add_heading("3.8 ハイパーパラメータ感度", level=2)
    add_body(
        doc,
        "自由rollout前のteacher-forced calibrationとして、rerank pool 64–1,024、decision strength 1–16の25条件を3 seedで評価した。"
        "pool 64では最大誤差が1%を超える条件が存在したが、pool 128以上では全条件がteacher-forced最大1%未満であった。"
        "ただしteacher-forced最良条件だけでは自由rollout最大誤差を保証できず、最終的には訪問状態のexact再ラベルが必要であった。"
    )
    add_figure(
        doc,
        "Figure_9_hyperparameter_calibration.png",
        6.25,
        "図9　teacher-forced状態におけるハイパーパラメータ感度。（a）3 seed平均PHA積分誤差。（b）3 seed最大誤差。セルは相対誤差（%）。この評価はCPU exact状態上の局所評価であり、GPU自由rolloutの終点誤差とは区別される。",
        "Heatmaps of teacher-forced PHA error across rerank-pool and decision-strength settings.",
    )

    doc.add_heading("3.9 適格性と安全性", level=2)
    add_body(
        doc,
        "最終モデルはPHA最大誤差基準の86.1%、菌体量誤差基準の5.4%であり、いずれも許容上限内であった。"
        "validationの600要求は全件GPU有効解で、4環境smoke testでも512要求を全件処理した。5件は候補拡張でGPU上回復し、"
        "オンラインCPU LP呼出しは0件であった。"
    )
    add_figure(
        doc,
        "Figure_10_qualification_and_safety.png",
        6.25,
        "図10　事前定義した適格基準と実行時安全性。（a）最大PHA終点誤差および最大菌体量誤差を各許容上限で規格化した値。破線1.0以下を合格とする。（b）validationおよびsmoke testで処理したGPU要求、GPU再試行回復、CPU LP fallback件数。",
        "Normalized acceptance margins and counts of valid GPU requests, recovered retries, and CPU fallbacks.",
    )
    criteria_table = doc.add_table(rows=1, cols=4)
    fill_table(
        criteria_table,
        ["判定項目", "事前基準", "実測値", "判定"],
        [
            ["PHA終点相対誤差", "最大≤1.0%", f"最大 {100*validation['summary']['pha_relative_error_max']:.3f}%", "適合"],
            ["GPU解受理率", "≥90%", f"{100*validation['summary']['acceptance_rate_min']:.1f}%", "適合"],
            ["菌体量終点誤差", "最大≤0.01 g/L", f"{validation['summary']['max_biomass_absolute_error_g_l']:.6f} g/L", "適合"],
            ["速度", "平均speed-up≥1", f"{validation['summary']['speedup_mean']:.2f}×", "適合"],
            ["Online CPU LP", "0（GPU-only目標）", "0/600", "適合"],
        ],
        [2600, 2200, 2700, 1860],
        font_size=8.0,
    )
    doc.add_paragraph("表4　最終モデルの事前定義適格条件と判定。", style="Caption")

    doc.add_heading("4. 新規性の評価", level=1)
    add_body(
        doc,
        "本実装の新規性は、構成要素単体ではなく統合方法にある。共同培養dFBA、RL制御、ニューラル・機構論GEM、GPU batch optimizationは既知である。"
        "一方、厳密共同FBAアンカーを凸包基底としてSv = 0を構成上保持し、liveな共有培地制約をGPU QPで投影し、"
        "PPOの複数環境を単一CUDA serviceへ集約し、境界状態のみアンカー凸包を適応拡張する組合せは、検索した文献では同一報告を確認できなかった。"
    )
    novelty = doc.add_table(rows=1, cols=4)
    fill_table(
        novelty,
        ["研究領域", "先行研究で確立", "本実装の追加", "新規性判断"],
        [
            ["共同培養dFBA", "複数種の動的基質利用[2]", "3 GEM共有培地＋流加RL", "応用統合"],
            ["共培養RL", "population制御[4]、GEM agent[5]", "PPO multi-envをGPU FBAへ直結", "統合候補"],
            ["Neural–mechanistic GEM", "GEMを学習構造へ組込む[6]", "厳密flux anchorの凸包", "方法候補"],
            ["GPU最適化", "batch LP/QP、PDHG、cuOpt[7–9]", "shared-medium FBA特化QP", "領域特化"],
            ["失敗処理", "solver fallback/warm start", "128→2,048 anchorのGPU-only retry", "実装候補"],
        ],
        [1850, 2750, 3000, 1760],
        font_size=7.8,
    )
    doc.add_paragraph("表5　先行研究に対する本実装の位置付け。", style="Caption")
    add_body(
        doc,
        "ただし、“世界初”を主張するには不十分である。本調査は系統的レビューや特許調査ではなく、コード実装も単一研究系での検証である。"
        "論文上は『既存手法の新規な統合』『exact-feasible anchor convex-hull projection for multi-GEM dFBA』と表現し、"
        "査読前に検索式、データベース、期間を明示したsystematic searchと、各構成要素を除くablation studyを追加することが妥当である。"
    )

    doc.add_heading("5. 考察", level=1)
    add_body(
        doc,
        "高速化の主要因は、6,585-flux LPの疎行列構築と三段階HiGHS求解をonline loopから除き、"
        "GPU常駐辞書のmatrix operationとbatch projectionへ置換した点にある。batch size増加に伴う環境当たりlatency低下は、"
        "GPU利用には単一環境よりも多数環境の並列実行が適するという設計仮説を支持する。"
    )
    add_body(
        doc,
        f"ただし、観測された{validation['summary']['speedup_mean']:.2f}倍はCPU LPと近似GPU系のsystem-level比較であり、同一最適化問題を異なるhardwareで厳密に解くsolver benchmarkではない。"
        "GPU系は辞書離散化とニューラル順位付けを含み、CPU三段階目的を完全には再現しない。"
        "さらにdFBAの時間依存性、Python状態更新、worker通信はCPUに残るため、kernel throughputの7.82倍がそのまま学習全体へ反映されるわけではない。"
    )
    add_body(
        doc,
        f"PHA誤差の改善は、単純な候補数増加ではなく、GPU自身が訪問する状態をexact解で被覆したことに由来する。"
        f"720アンカー（baseline比2.14%）の追加で平均誤差を{pha_mean_reduction:.1f}%低減できたことは、"
        "辞書全体を一様に拡張するよりもrollout分布へ焦点を当てたactiveな追加が効率的であることを示唆する。"
        "一方、n = 5であり、異なるpolicyが到達する領域では同じ精度を保証しない。"
    )
    add_body(
        doc,
        "複数GPUではVRAMを共有せず、seed、agentまたは環境群を独立shardとして割り当てることが適切である。"
        "本稿ではRTX 4060 Laptop GPU 1基のみを実測しており、RTX PRO 4000複数基への性能外挿は実測結果ではない。"
        "3 GPUの主張には、1/2/3 GPUで同一総transition数を用いたstrong scalingおよびGPU当たり固定環境数を用いたweak scalingが必要である。"
    )

    doc.add_heading("6. 限界と今後の検証", level=1)
    limitations = [
        "PHA最大相対誤差0.861%は本benchmarkの1%基準を満たすが、学習済みPPOが生成する異なるaction分布について再validationが必要である。",
        "n = 5は性能再現性の初期確認であり、異なる培地、初期菌体量、gene knockout、policy分布、長期軌道を含む外部validationが必要である。",
        "GPU kernel scalingは3反復であり、energy consumption、GPU utilization、host overheadを含むprofilingが未完了である。",
        "CPU–GPU比較は近似手法と厳密手法の比較である。cuOpt barrier/PDLP等による同一LPのGPU厳密解を別baselineとして追加すべきである[9]。",
        "PHA値はGEM/dFBA内のモデル出力であり、実培養でのPHA定量、菌体量、基質濃度による外部妥当性確認を要する。",
        "新規性評価は探索的であり、systematic literature review、特許調査、公開benchmarkおよびablationが必要である。",
    ]
    for index, item in enumerate(limitations, 1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run(item), "IPAexMincho", 10)

    doc.add_heading("7. 結論", level=1)
    add_body(
        doc,
        "三種GEM共同dFBAの反復CPU LPを、厳密flux anchorのCUDA順位付けと凸包QP投影へ置換し、PPO複数環境を単一GPU serviceへ集約した。"
        f"RTX 4060 Laptop GPU上で24時間軌道を平均{validation['summary']['speedup_mean']:.2f}倍高速化し、batch 64でkernel throughputをbatch 1比7.82倍へ拡張した。"
        "適応的GPU再試行により4環境smoke testの512要求を全件有効解とし、online CPU LP呼出しを0にできた。"
        f"さらにDAgger型exact再ラベルでPHA終点誤差を平均{100*baseline['summary']['pha_relative_error_mean']:.2f}%から"
        f"{100*validation['summary']['pha_relative_error_mean']:.2f}%、最大{100*baseline['summary']['pha_relative_error_max']:.2f}%から"
        f"{100*validation['summary']['pha_relative_error_max']:.2f}%へ低減し、全held-out seedで最大1%基準を満たした。"
        "本構成は高スループット探索用として有用であり、exact-feasible anchor凸包、shared-medium QP、rollout状態の機構論的再ラベルを統合した点に方法的新規性候補がある。"
        "ただし最終科学評価では、学習済みpolicyのCPU exact再評価と実培養validationを併用する必要がある。"
    )

    doc.add_heading("データおよびコードの可用性", level=1)
    add_body(
        doc,
        "主要実装はsrc/gpu_batch_qp.py、src/cooperative_gpu_service.py、src/cooperative_neural_surrogate.py、src/community_solver.pyにある。"
        "図の一次データはbaseline/final validation manifest、DAgger反復report、PHA calibration JSON、GPU profile、"
        "GPU batch scaling JSON、parallel training smoke JSONであり、図生成コードはscripts/create_cpu_gpu_paper_figures.pyに保存した。"
        "最終実装では全pytest 88件が合格した。"
    )

    doc.add_heading("参考文献", level=1)
    references = [
        ("Orth JD, Thiele I, Palsson BØ. What is flux balance analysis? Nat Biotechnol. 2010;28:245–248.", "https://doi.org/10.1038/nbt.1614"),
        ("Hanly TJ, Henson MA. Dynamic flux balance modeling of microbial co-cultures. Biotechnol Bioeng. 2011;108:376–385.", "https://doi.org/10.1002/bit.23007"),
        ("Höffner K, Harwood SM, Barton PI. A reliable simulator for dynamic flux balance analysis. Biotechnol Bioeng. 2013;110:792–802.", "https://doi.org/10.1002/bit.24748"),
        ("Treloar NJ, Fedorec AJH, Ingalls B, Barnes CP. Deep reinforcement learning for the control of microbial co-cultures in bioreactors. PLoS Comput Biol. 2020;16:e1007783.", "https://doi.org/10.1371/journal.pcbi.1007783"),
        ("Ghadermazi P, Chan SHJ. Microbial interactions from a new perspective: reinforcement learning reveals new insights into microbiome evolution. Bioinformatics. 2024;40:btae003.", "https://doi.org/10.1093/bioinformatics/btae003"),
        ("Faure L, Mollet B, Liebermeister W, Faulon J-L, et al. A neural-mechanistic hybrid approach improving the predictive power of genome-scale metabolic models. Nat Commun. 2023;14:4669.", "https://doi.org/10.1038/s41467-023-40380-0"),
        ("Gurung A, Ray R. Solving Batched Linear Programs on GPU and Multicore CPU. arXiv:1609.08114. 2016.", "https://arxiv.org/abs/1609.08114"),
        ("Amos B, Kolter JZ. OptNet: Differentiable Optimization as a Layer in Neural Networks. Proc ICML. 2017;70:136–145.", "https://proceedings.mlr.press/v70/amos17a.html"),
        ("NVIDIA. cuOpt LP/QP Features: PDLP, barrier, quadratic programming, and batch mode. User Guide 26.04.", "https://docs.nvidia.com/cuopt/user-guide/latest/lp-qp-features.html"),
        ("Hallmann N, Guerra-Cornejo C, Burgess K, Merzbacher C, Oyarzún DA. Multiobjective Design of Growth Media with Genome-Scale Metabolic Models and Bayesian Optimization. CSBJ. 2026;35:0072.", "https://doi.org/10.34133/csbj.0072"),
        ("Ross S, Gordon G, Bagnell D. A Reduction of Imitation Learning and Structured Prediction to No-Regret Online Learning. Proc AISTATS. 2011;15:627–635.", "https://proceedings.mlr.press/v15/ross11a.html"),
        ("Bengio S, Vinyals O, Jaitly N, Shazeer N. Scheduled Sampling for Sequence Prediction with Recurrent Neural Networks. Adv Neural Inf Process Syst. 2015;28.", "https://proceedings.neurips.cc/paper/2015/hash/e995f98d56967d946471af29d7bf99f1-Abstract.html"),
    ]
    for index, (citation, url) in enumerate(references, 1):
        add_reference(doc, index, citation, url)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    set_cell_shading_like_paragraph(note, CAUTION_FILL)
    set_font(note.add_run("査読前注記　"), "IPAexGothic", 9, bold=True, color=RGBColor(154, 84, 0))
    set_font(
        note.add_run("本稿は実装・ベンチマークの技術原稿であり、生物学的実験validationおよびsystematic novelty reviewの完了を意味しない。"),
        "IPAexMincho",
        9,
        color=INK,
    )

    doc.save(DOCX)
    print(DOCX)


def set_cell_shading_like_paragraph(paragraph, fill: str):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "D55E00")
    borders.append(left)
    ppr.append(borders)


if __name__ == "__main__":
    main()
