from pathlib import Path
from docx import Document
from docx.shared import RGBColor
ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/revised_20260922'
OUT.mkdir(exist_ok=True)
notice = '2026年9月22日改訂　研究室ではRoxA/RoxBの異種宿主発現に成功していない。今回の酸素要求量測定はNS21由来クルードを用いる。文献の精製酵素条件は参考情報であり、研究室での実施済み条件を意味しない。'
details = [
 '試料はNS21培養由来の無細胞クルードを第一候補とする。培養上清、菌体関連画分または破砕抽出液のどこに活性が回収されるかを予備確認し、由来画分を明記する。残存菌体の有無と分離操作前後の活性回収を確認する。無細胞上清と全培養液・休止細胞を同一試料として扱わない。',
 '各DO条件で、活性クルード＋ゴム、同じクルード＋ゴムなし、同じ工程で処理した未接種培地＋ゴム、同培地＋ゴムなしを測定する。加熱失活クルードの基質あり・なしも補助対照とするが、加熱によるマトリクス変化があるため、それだけでRox特異性を保証しない。基質溶媒・添加体積は対応させる。',
 'DO低下と並行して、基質添加前後の切断生成物をHPLC等で確認する。背景補正した酸素消費をゴム添加に伴う見かけの消費として記録し、他の酸化酵素や残存呼吸の寄与が除けなければRoxA/RoxB固有活性と断定しない。ODTDやオリゴマー分布だけで各酵素の寄与率を一意に分離しない。',
 'クルード添加量、総タンパク質量、元培養液換算体積、培養条件・採取時点、濃縮倍率、保存履歴を記録する。主結果は酸素消費速度と生成物増加速度を実単位で報告し、総タンパク質当たりの値はクルード比活性と呼ぶ。精製酵素当たりの比活性、kcat、RoxAとRoxBの個別Kmへ換算しない。独立培養由来の複数ロットで再現性を確認する。',
 '基質量とクルード量を予備検討し、初期速度の測定可能域と量依存性を確認する。基質0.1%を自動的に飽和条件とみなさない。DOは実測濃度で記録し、温度・pH・撹拌・基質を固定する。低DO側を含む系列を設け、センサー応答と再曝気・漏入の影響を空試験で確認する。曝気停止だけで酸素移動をゼロと仮定しない。',
 '単一飽和式がデータに適合し、低DO域と高DO域の両方を測定できた場合のみ、クルード系の見かけの酸素半飽和定数 K0.5,app を不確かさ付きで推定する。適合しない場合は応答曲線または推定範囲を報告する。酸素消費量／生成物量の見かけの比と半飽和定数は異なる測定量である。',
 'モデルへは試料・DO・培養条件に対応する複合ゴム分解系の見かけの酸素応答として反映する。RoxA/RoxB別の定数や全菌種共通の定数を直接確定しない。全菌体でのゴムあり／なしのOUR差には分解産物の呼吸も含まれるため、酵素消費率や固定のpolymer_oxygen_fractionの直接測定とは扱わない。無細胞クルード、洗浄菌体、全培養系を別測定し、再構成の収支が検証できるまで分配は未同定とする。',
 '測定原理の参考はRöther et al. 2017, DOI 10.21769/BioProtoc.2188（酸素消費測定と切断生成物のHPLC分析）。ここに記したクルード対照・酸素濃度系列は本研究用の提案であり、同論文で実証済みのクルード手順とは位置づけない。'
]
def replace_section(doc, start, end, heading, body):
    ps=doc.paragraphs
    a=next(p for p in ps if p.text.startswith(start)); b=next(p for p in ps if p.text.startswith(end))
    node=a._p.getnext()
    while node is not None and node is not b._p:
        nxt=node.getnext();node.getparent().remove(node);node=nxt
    a.text=heading
    for text in body: b.insert_paragraph_before(text)
files=['OXYGEN_PARTITIONING_PROTOCOL_20260912.docx','LCP_ROX_ACTIVITY_ASSAY_PROTOCOL_20260911.docx','CALIBRATION_EXPERIMENT_PROTOCOL_20260911.docx']
for name in files:
    doc=Document(ROOT/'docs'/name)
    doc.paragraphs[1].insert_paragraph_before(notice)
    if name.startswith('OXYGEN'):
        replace_section(doc,'3.1','3.2','3.1 NS21クルードの酸素依存性測定',details)
        replace_section(doc,'3.3','4.','3.3 培養系全体の酸素消費との対応',[details[6],'全培養OURとクルード由来の速度は、試料体積・酵素回収率・DO・基質条件が異なるまま単純に差し引かない。酵素の発現・精製成功を本段階の開始条件としない。'])
        for p in doc.paragraphs:
            if '3.3節のOUR差分法による' in p.text: p.text='次の優先作業は、クルードの活性回収と背景消費の確認、DO応答測定、全菌体測定との対応づけである。単純差分から酸素分配率を確定しない。'
            if 'OUR差分法(3.3節)の解釈' in p.text: p.text=details[6]
    elif name.startswith('LCP_'):
        replace_section(doc,'3.2','3.3','3.2 NS21のクルードを用いたRox関連活性測定',details)
    else:
        replace_section(doc,'4.6','4.7','4.6 ゴム分解速度の測定',[notice,details[0],details[1],details[2],details[3]])
        replace_section(doc,'4.7','4.8','4.7 酸素応答の測定',[details[4],details[5],details[6]])
    for p in doc.paragraphs:
        if p.style.name.startswith(('Heading','Title')):
            for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
    dest=OUT/name.replace('.docx','_CRUDE_REV20260922.docx')
    doc.save(dest);print(dest)
